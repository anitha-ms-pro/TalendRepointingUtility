"""
Talend Unique Components Finder
Scans all .item files in a repository and lists all unique Talend components used.
Compares against migration config to identify if components are configured for migration.
"""

import os
import xml.etree.ElementTree as ET
from pathlib import Path
from collections import defaultdict
import argparse
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("Warning: openpyxl not installed. Excel output will not be available.")
    print("Install with: pip install openpyxl")

# Import the config to check component mappings
try:
    from config import COMPONENT_REPLACEMENTS, COMPONENTS_TO_REMOVE
    CONFIG_AVAILABLE = True
except ImportError:
    CONFIG_AVAILABLE = False
    print("Warning: config.py not found. Migration status will not be available.")


def find_item_files(repo_path):
    """Recursively find all .item files in the repository."""
    item_files = []
    for root, dirs, files in os.walk(repo_path):
        for file in files:
            if file.endswith('.item'):
                item_files.append(os.path.join(root, file))
    return item_files


def extract_components_from_item(item_file_path):
    """Extract all component names from a Talend .item file."""
    components = set()

    try:
        tree = ET.parse(item_file_path)
        root = tree.getroot()

        # Find all elements with componentName attribute
        for elem in root.iter():
            if 'componentName' in elem.attrib:
                component_name = elem.attrib['componentName']
                components.add(component_name)

        # Also check for xmi:type attributes which contain component types
        for elem in root.iter():
            xmi_type = elem.attrib.get('{http://www.omg.org/XMI}type', '')
            if xmi_type and ':' in xmi_type:
                component_type = xmi_type.split(':')[-1]
                # Filter out non-component types
                if not component_type.endswith('Connection') and \
                   not component_type.endswith('Item') and \
                   component_type not in ['ProcessType', 'NodeType', 'ConnectionType',
                                          'ElementParameterType', 'RoutinesParameterType',
                                          'MetadataType', 'ColumnType']:
                    components.add(component_type)

    except ET.ParseError as e:
        print(f"Warning: Failed to parse {item_file_path}: {e}")
    except Exception as e:
        print(f"Warning: Error processing {item_file_path}: {e}")

    return components


def get_migration_status(component_name):
    """
    Check if a component is configured for migration.
    Returns: (status, target_component, notes)
    """
    if not CONFIG_AVAILABLE:
        return "UNKNOWN", "", "Config not available"

    if component_name in COMPONENT_REPLACEMENTS:
        target = COMPONENT_REPLACEMENTS[component_name]
        if target is None:
            return "REMOVE", "", "Component will be removed (no GCP equivalent)"
        else:
            return "MIGRATE", target, f"Will be replaced with {target}"
    else:
        # Component not in config - might be okay as-is or needs review
        return "NOT_IN_CONFIG", "", "Component not in migration config - needs review"


def find_all_components(repo_path, verbose=False):
    """Find all unique components across all .item files."""
    print(f"Scanning repository: {repo_path}")

    item_files = find_item_files(repo_path)
    print(f"Found {len(item_files)} .item files")

    all_components = set()
    component_locations = defaultdict(list)  # Track where each component is used

    for item_file in item_files:
        if verbose:
            print(f"Processing: {item_file}")

        components = extract_components_from_item(item_file)
        all_components.update(components)

        # Track locations for detailed output
        for component in components:
            relative_path = os.path.relpath(item_file, repo_path)
            component_locations[component].append(relative_path)

    return all_components, component_locations, len(item_files)


def write_excel_report(components, locations, total_files, output_path, repo_path):
    """Write the component report to an Excel file."""
    if not EXCEL_AVAILABLE:
        print("Error: openpyxl is not installed. Cannot create Excel file.")
        return False

    wb = openpyxl.Workbook()

    # Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"

    # Header styling
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)

    # Summary information
    ws_summary['A1'] = "Talend Component Analysis Report"
    ws_summary['A1'].font = Font(bold=True, size=16)
    ws_summary.merge_cells('A1:D1')

    ws_summary['A3'] = "Repository Path:"
    ws_summary['B3'] = os.path.abspath(repo_path)
    ws_summary['A4'] = "Total .item Files:"
    ws_summary['B4'] = total_files
    ws_summary['A5'] = "Unique Components:"
    ws_summary['B5'] = len(components)

    # Migration status summary
    ws_summary['A7'] = "Migration Status Summary"
    ws_summary['A7'].font = Font(bold=True, size=14)

    status_counts = {"MIGRATE": 0, "REMOVE": 0, "NOT_IN_CONFIG": 0, "UNKNOWN": 0}
    for comp in components:
        status, _, _ = get_migration_status(comp)
        status_counts[status] += 1

    ws_summary['A8'] = "Components to Migrate:"
    ws_summary['B8'] = status_counts["MIGRATE"]
    ws_summary['A9'] = "Components to Remove:"
    ws_summary['B9'] = status_counts["REMOVE"]
    ws_summary['A10'] = "Components Not in Config:"
    ws_summary['B10'] = status_counts["NOT_IN_CONFIG"]
    if status_counts["UNKNOWN"] > 0:
        ws_summary['A11'] = "Unknown Status:"
        ws_summary['B11'] = status_counts["UNKNOWN"]

    # Components Detail Sheet
    ws_detail = wb.create_sheet("Components Detail")

    # Headers
    headers = ["#", "Component Name", "Usage Count", "Migration Status", "Target Component", "Notes"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws_detail.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sort components by usage count (descending)
    sorted_components = sorted(components, key=lambda c: len(locations[c]), reverse=True)

    # Data rows
    for idx, component in enumerate(sorted_components, 1):
        usage_count = len(locations[component])
        status, target, notes = get_migration_status(component)

        row = idx + 1
        ws_detail.cell(row=row, column=1).value = idx
        ws_detail.cell(row=row, column=2).value = component
        ws_detail.cell(row=row, column=3).value = usage_count
        ws_detail.cell(row=row, column=4).value = status
        ws_detail.cell(row=row, column=5).value = target
        ws_detail.cell(row=row, column=6).value = notes

        # Color coding based on status
        status_cell = ws_detail.cell(row=row, column=4)
        if status == "MIGRATE":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Light green
        elif status == "REMOVE":
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Light red
        elif status == "NOT_IN_CONFIG":
            status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Light yellow

    # Adjust column widths
    ws_detail.column_dimensions['A'].width = 5
    ws_detail.column_dimensions['B'].width = 30
    ws_detail.column_dimensions['C'].width = 12
    ws_detail.column_dimensions['D'].width = 18
    ws_detail.column_dimensions['E'].width = 25
    ws_detail.column_dimensions['F'].width = 50

    # File Locations Sheet
    ws_locations = wb.create_sheet("Component Locations")

    # Headers
    ws_locations.cell(row=1, column=1).value = "Component Name"
    ws_locations.cell(row=1, column=2).value = "File Path"
    ws_locations.cell(row=1, column=1).font = header_font
    ws_locations.cell(row=1, column=1).fill = header_fill
    ws_locations.cell(row=1, column=2).font = header_font
    ws_locations.cell(row=1, column=2).fill = header_fill

    # Data rows
    row_idx = 2
    for component in sorted_components:
        for file_path in sorted(locations[component]):
            ws_locations.cell(row=row_idx, column=1).value = component
            ws_locations.cell(row=row_idx, column=2).value = file_path
            row_idx += 1

    ws_locations.column_dimensions['A'].width = 30
    ws_locations.column_dimensions['B'].width = 80

    # Save workbook
    wb.save(output_path)
    print(f"\n[OK] Excel report saved to: {output_path}")
    return True


def write_text_report(components, locations, total_files, output_path, repo_path):
    """Write the component report to a text file."""
    output_lines = []
    output_lines.append("=" * 100)
    output_lines.append("TALEND UNIQUE COMPONENTS REPORT")
    output_lines.append("=" * 100)
    output_lines.append(f"Repository: {os.path.abspath(repo_path)}")
    output_lines.append(f"Total .item files scanned: {total_files}")
    output_lines.append(f"Unique components found: {len(components)}")
    output_lines.append("=" * 100)
    output_lines.append("")

    # Sort components by usage count
    sorted_components = sorted(components, key=lambda c: len(locations[c]), reverse=True)

    # Component list with migration status
    output_lines.append("COMPONENT LIST WITH MIGRATION STATUS:")
    output_lines.append("-" * 100)
    output_lines.append(f"{'#':<5} {'Component Name':<35} {'Usage':<8} {'Status':<18} {'Target Component':<25}")
    output_lines.append("-" * 100)

    for idx, component in enumerate(sorted_components, 1):
        usage_count = len(locations[component])
        status, target, notes = get_migration_status(component)

        status_symbol = {
            "MIGRATE": "[+]",
            "REMOVE": "[-]",
            "NOT_IN_CONFIG": "[?]",
            "UNKNOWN": "[?]"
        }.get(status, "[?]")

        output_lines.append(
            f"{idx:<5} {component:<35} {usage_count:<8} "
            f"{status_symbol} {status:<16} {target:<25}"
        )

    output_lines.append("")
    output_lines.append("=" * 100)
    output_lines.append("Legend:")
    output_lines.append("  [+] MIGRATE         - Component will be migrated to GCP equivalent")
    output_lines.append("  [-] REMOVE          - Component will be removed (no GCP equivalent)")
    output_lines.append("  [?] NOT_IN_CONFIG   - Component not in migration config (needs review)")
    output_lines.append("=" * 100)

    output_text = "\n".join(output_lines)

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(output_text)

    print(f"\n[OK] Text report saved to: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description='Find all unique Talend components used in a repository and check migration status'
    )
    parser.add_argument(
        'repo_path',
        help='Path to the Talend repository'
    )
    parser.add_argument(
        '-o', '--output',
        required=True,
        help='Output Excel file path (e.g., components_report.xlsx)'
    )
    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='Show detailed output during processing'
    )
    parser.add_argument(
        '--text-report',
        help='Also generate a text report at this path'
    )

    args = parser.parse_args()

    # Validate repository path
    if not os.path.exists(args.repo_path):
        print(f"Error: Repository path does not exist: {args.repo_path}")
        return 1

    # Validate output path
    if not args.output.endswith('.xlsx'):
        print("Error: Output file must have .xlsx extension")
        return 1

    # Find all components
    components, locations, total_files = find_all_components(args.repo_path, args.verbose)

    if not components:
        print("No components found in the repository.")
        return 0

    # Write Excel report
    success = write_excel_report(components, locations, total_files, args.output, args.repo_path)

    if not success:
        print("Failed to create Excel report. Check that openpyxl is installed.")
        return 1

    # Optionally write text report
    if args.text_report:
        write_text_report(components, locations, total_files, args.text_report, args.repo_path)

    # Print summary
    print("\n" + "=" * 100)
    print("SUMMARY")
    print("=" * 100)
    print(f"Repository: {args.repo_path}")
    print(f"Total .item files: {total_files}")
    print(f"Unique components: {len(components)}")

    if CONFIG_AVAILABLE:
        status_counts = {"MIGRATE": 0, "REMOVE": 0, "NOT_IN_CONFIG": 0}
        for comp in components:
            status, _, _ = get_migration_status(comp)
            status_counts[status] += 1

        print(f"\nMigration Status:")
        print(f"  [+] Components to migrate: {status_counts['MIGRATE']}")
        print(f"  [-] Components to remove: {status_counts['REMOVE']}")
        print(f"  [?] Not in config: {status_counts['NOT_IN_CONFIG']}")

    print("=" * 100)

    return 0


if __name__ == '__main__':
    exit(main())
